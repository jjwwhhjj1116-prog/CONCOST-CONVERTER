"""
표 및 문서 추출 엔진 (Table & Document Engine)
도면/내역서/일람표 PDF에서 격자 선분을 고속 복원하고 텍스트를 셀 단위로 정확하게 매핑하여 Excel로 변환합니다.
(CONCOST 검증 알고리즘 + 얇은 사각형 중심선 치환 및 공선 병합 최적화)
"""
from dataclasses import dataclass
from pathlib import Path
from statistics import median
from typing import Any, Dict, List, Optional, Tuple

import fitz  # PyMuPDF
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from shapely.geometry import LineString, Point, box
from shapely.ops import polygonize, unary_union
from shapely.strtree import STRtree


@dataclass
class TableCell:
    """복원된 개별 표 셀"""
    row: int
    col: int
    text: str
    bbox: Tuple[float, float, float, float]


@dataclass
class RecoveredTable:
    """복원된 표 구조"""
    table_index: int
    rows_count: int
    cols_count: int
    grid: List[List[str]]  # [행][열] 텍스트 매트릭스


class TableEngine:
    """PDF 표 복원 및 Excel 변환 엔진"""

    @staticmethod
    def extract_lines_and_texts_from_page(
        page: fitz.Page,
    ) -> Tuple[List[Tuple[Tuple[float, float], Tuple[float, float]]], List[Dict[str, Any]]]:
        """단일 PyMuPDF 페이지 객체에서 격자 후보 선분과 텍스트 단어를 추출합니다."""
        lines: List[Tuple[Tuple[float, float], Tuple[float, float]]] = []
        texts: List[Dict[str, Any]] = []

        for draw in page.get_drawings():
            items = draw.get("items", [])
            # 글리프 윤곽선이나 복잡한 해치 경로(수십 개 곡선으로 구성된 단일 드로잉)는 표 테두리가 아니므로 제외
            if len(items) > 60 and any(it[0] == "c" for it in items):
                continue
            for item in items:
                op = item[0]
                if op == "l":  # line
                    p1, p2 = item[1], item[2]
                    lines.append(((float(p1.x), float(p1.y)), (float(p2.x), float(p2.y))))
                elif op == "re":  # rect
                    r = item[1]
                    w = abs(float(r.x1) - float(r.x0))
                    h = abs(float(r.y1) - float(r.y0))
                    # CAD/Office PDF에서 표 선을 얇은 채움 사각형(두께 <= 3.5pt)으로 출력한 경우 중심선 1개로 치환
                    if w <= 3.5 and h > 6.0:
                        cx = (float(r.x0) + float(r.x1)) / 2.0
                        lines.append(((cx, float(r.y0)), (cx, float(r.y1))))
                    elif h <= 3.5 and w > 6.0:
                        cy = (float(r.y0) + float(r.y1)) / 2.0
                        lines.append(((float(r.x0), cy), (float(r.x1), cy)))
                    elif w > 3.5 and h > 3.5:
                        lines.append(((float(r.x0), float(r.y0)), (float(r.x1), float(r.y0))))
                        lines.append(((float(r.x1), float(r.y0)), (float(r.x1), float(r.y1))))
                        lines.append(((float(r.x1), float(r.y1)), (float(r.x0), float(r.y1))))
                        lines.append(((float(r.x0), float(r.y1)), (float(r.x0), float(r.y0))))

        for word in page.get_text("words"):
            x0, y0, x1, y1, text = word[:5]
            texts.append({
                "text": str(text),
                "bbox": (float(x0), float(y0), float(x1), float(y1)),
                "anchor": ((float(x0) + float(x1)) / 2.0, (float(y0) + float(y1)) / 2.0),
            })

        return lines, texts

    @classmethod
    def extract_lines_and_texts(
        cls, pdf_path: Path, page_index: int = 0
    ) -> Tuple[List[Tuple[Tuple[float, float], Tuple[float, float]]], List[Dict[str, Any]]]:
        """PDF 파일의 특정 페이지에서 선분 좌표 쌍과 텍스트 바운딩 박스를 추출합니다."""
        with fitz.open(pdf_path) as doc:
            if page_index >= len(doc):
                return [], []
            return cls.extract_lines_and_texts_from_page(doc[page_index])

    @classmethod
    def recover_tables(
        cls,
        lines: List[Tuple[Tuple[float, float], Tuple[float, float]]],
        texts: List[Dict[str, Any]],
        tolerance: Optional[float] = None,
    ) -> List[RecoveredTable]:
        """선분 격자를 클러스터링·병합하고 닫힌 사각형 셀을 복원하여 텍스트를 배치합니다."""
        if not texts:
            return []

        heights = [abs(t["bbox"][3] - t["bbox"][1]) for t in texts if t["bbox"][3] > t["bbox"][1]]
        tol = tolerance or max(1.2, (median(heights) * 0.22 if heights else 1.5))
        min_seg_len = max(tol * 3.0, 8.0)

        segments: List[Tuple[str, float, float, float]] = []
        xs: List[float] = []
        ys: List[float] = []
        for a, b in lines:
            x0, y0 = float(a[0]), float(a[1])
            x1, y1 = float(b[0]), float(b[1])
            if abs(y1 - y0) <= tol and abs(x1 - x0) >= min_seg_len:
                y = (y0 + y1) / 2.0
                segments.append(("h", min(x0, x1), max(x0, x1), y))
                xs.extend((x0, x1))
                ys.append(y)
            elif abs(x1 - x0) <= tol and abs(y1 - y0) >= min_seg_len:
                x = (x0 + x1) / 2.0
                segments.append(("v", min(y0, y1), max(y0, y1), x))
                ys.extend((y0, y1))
                xs.append(x)

        if not segments:
            return cls._recover_text_coordinate_grid(texts, tol)

        def cluster(values: List[float]) -> Dict[float, float]:
            res: Dict[float, float] = {}
            grp: List[float] = []
            for v in sorted(set(values)):
                if grp and v - grp[0] > tol * 1.5:
                    c = round(sum(grp) / len(grp), 2)
                    res.update((x, c) for x in grp)
                    grp = []
                grp.append(v)
            if grp:
                c = round(sum(grp) / len(grp), 2)
                res.update((x, c) for x in grp)
            return res

        xm, ym = cluster(xs), cluster(ys)

        # 동일 행(y) 또는 동일 열(x) 상의 겹치는 선분 구간을 병합하여 Shapely 연산량을 대폭 축소
        h_by_y: Dict[float, List[Tuple[float, float]]] = {}
        v_by_x: Dict[float, List[Tuple[float, float]]] = {}
        for axis, lo, hi, k in segments:
            if axis == "h":
                c_lo, c_hi, c_y = xm[lo], xm[hi], ym[k]
                if c_hi - c_lo >= min_seg_len:
                    h_by_y.setdefault(c_y, []).append((c_lo, c_hi))
            else:
                c_lo, c_hi, c_x = ym[lo], ym[hi], xm[k]
                if c_hi - c_lo >= min_seg_len:
                    v_by_x.setdefault(c_x, []).append((c_lo, c_hi))

        def merge_intervals(intervals: List[Tuple[float, float]]) -> List[Tuple[float, float]]:
            if not intervals:
                return []
            s_int = sorted(intervals)
            merged = [s_int[0]]
            for cur_lo, cur_hi in s_int[1:]:
                prev_lo, prev_hi = merged[-1]
                if cur_lo <= prev_hi + tol * 2.0:
                    merged[-1] = (prev_lo, max(prev_hi, cur_hi))
                else:
                    merged.append((cur_lo, cur_hi))
            return merged

        edges: List[LineString] = []
        for cy, ivs in h_by_y.items():
            for lo, hi in merge_intervals(ivs):
                edges.append(LineString([(lo, cy), (hi, cy)]))
        for cx, ivs in v_by_x.items():
            for lo, hi in merge_intervals(ivs):
                edges.append(LineString([(cx, lo), (cx, hi)]))

        if len(edges) < 4:
            return cls._recover_text_coordinate_grid(texts, tol)

        cells = []
        for poly in polygonize(unary_union(edges)):
            b = poly.bounds
            w = b[2] - b[0]
            h = b[3] - b[1]
            if w > tol * 2.0 and h > tol * 1.5:
                cells.append(box(*b))

        if not cells:
            return cls._recover_text_coordinate_grid(texts, tol)

        # R-Tree 공간 인덱싱으로 텍스트 중심점을 포함하는 가장 작은 셀에 할당
        tree = STRtree(cells)
        assigned: Dict[int, List[Dict[str, Any]]] = {i: [] for i in range(len(cells))}
        cell_areas = [c.area for c in cells]

        for t in texts:
            anchor = t.get("anchor")
            if not anchor:
                continue
            pt = Point(anchor[0], anchor[1])
            hits = tree.query(pt, predicate="within")
            if len(hits) == 0:
                hits = tree.query(box(*t["bbox"]), predicate="intersects")
            if len(hits) > 0:
                best_cell = min((int(h_idx) for h_idx in hits), key=lambda idx: cell_areas[idx])
                assigned[best_cell].append(t)

        # 행(Y) 클러스터링 및 열(X) 정렬
        row_tol = max(4.0, tol * 2.5)
        sorted_cells = sorted(range(len(cells)), key=lambda i: (cells[i].bounds[1], cells[i].bounds[0]))

        row_groups: List[Tuple[float, List[int]]] = []
        for c_idx in sorted_cells:
            y_top = cells[c_idx].bounds[1]
            if row_groups and abs(y_top - row_groups[-1][0]) <= row_tol:
                row_groups[-1][1].append(c_idx)
            else:
                row_groups.append((y_top, [c_idx]))

        grid: List[List[str]] = []
        for _, r_cells in row_groups:
            r_cells_sorted = sorted(r_cells, key=lambda i: cells[i].bounds[0])
            row_texts: List[str] = []
            for c_idx in r_cells_sorted:
                words_in_cell = sorted(
                    assigned[c_idx],
                    key=lambda item: (round(item["bbox"][1] / 3.0), item["bbox"][0]),
                )
                row_texts.append(" ".join(w["text"] for w in words_in_cell).strip())
            if any(row_texts):
                grid.append(row_texts)

        if not grid:
            return cls._recover_text_coordinate_grid(texts, tol)

        return [
            RecoveredTable(
                table_index=0,
                rows_count=len(grid),
                cols_count=max((len(r) for r in grid), default=0),
                grid=grid,
            )
        ]

    @staticmethod
    def _recover_text_coordinate_grid(texts: List[Dict[str, Any]], tol: float) -> List[RecoveredTable]:
        """선분이 없는 PDF 표의 경우 텍스트 Y/X 좌표를 클러스터링하여 행·열 격자로 복원합니다."""
        if not texts:
            return []
        row_tol = max(4.0, tol * 2.0)
        sorted_t = sorted(texts, key=lambda t: (t["bbox"][1], t["bbox"][0]))
        rows: List[Tuple[float, List[Dict[str, Any]]]] = []
        for t in sorted_t:
            y0 = t["bbox"][1]
            if rows and abs(y0 - rows[-1][0]) <= row_tol:
                rows[-1][1].append(t)
            else:
                rows.append((y0, [t]))

        grid: List[List[str]] = []
        for _, r_items in rows:
            r_sorted = sorted(r_items, key=lambda x: x["bbox"][0])
            cells: List[str] = []
            prev_x1: Optional[float] = None
            for item in r_sorted:
                if prev_x1 is not None and (item["bbox"][0] - prev_x1) < 12.0 and cells:
                    cells[-1] = f"{cells[-1]} {item['text']}".strip()
                else:
                    cells.append(item["text"])
                prev_x1 = item["bbox"][2]
            if any(cells):
                grid.append(cells)

        return [
            RecoveredTable(
                table_index=0,
                rows_count=len(grid),
                cols_count=max((len(r) for r in grid), default=0),
                grid=grid,
            )
        ]

    @classmethod
    def convert_pdf_to_excel(cls, pdf_path: Path, excel_path: Path, page_index: int = 0) -> bool:
        """PDF 표를 복원하여 전문 Excel 서식 파일로 변환 저장합니다 (다중 페이지 지원)."""
        excel_path.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        default_ws = wb.active

        thin_side = Side(border_style="thin", color="A0A0A0")
        cell_border = Border(top=thin_side, left=thin_side, right=thin_side, bottom=thin_side)
        header_fill = PatternFill(start_color="F2F4F8", end_color="F2F4F8", fill_type="solid")
        header_font = Font(name="맑은 고딕", size=10, bold=True)
        body_font = Font(name="맑은 고딕", size=9)

        with fitz.open(pdf_path) as doc:
            total_pages = min(len(doc), 20)
            for p_idx in range(total_pages):
                ws = default_ws if p_idx == 0 else wb.create_sheet()
                ws.title = f"Page_{p_idx + 1}" if total_pages > 1 else "표 복원 결과"
                lines, texts = cls.extract_lines_and_texts_from_page(doc[p_idx])
                tables = cls.recover_tables(lines, texts)

                if not tables or not tables[0].grid:
                    for r_idx, t in enumerate(texts, start=1):
                        ws.cell(row=r_idx, column=1, value=t["text"]).font = body_font
                else:
                    table = tables[0]
                    for r_idx, row in enumerate(table.grid, start=1):
                        for c_idx, val in enumerate(row, start=1):
                            cell = ws.cell(row=r_idx, column=c_idx, value=val)
                            cell.border = cell_border
                            cell.alignment = Alignment(vertical="center", wrap_text=True)
                            if r_idx == 1:
                                cell.fill = header_fill
                                cell.font = header_font
                            else:
                                cell.font = body_font

                for col in ws.columns:
                    max_len = max((len(str(cell.value or "")) for cell in col), default=8)
                    col_letter = openpyxl.utils.get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = max(12, min(max_len * 1.5, 45))

        wb.save(str(excel_path))
        wb.close()
        return excel_path.exists() and excel_path.stat().st_size > 0

    @classmethod
    def convert_pdf_to_csv(cls, pdf_path: Path, csv_path: Path, page_index: int = 0) -> bool:
        """PDF 표를 복원하여 UTF-8 BOM CSV 파일로 변환 저장합니다."""
        import csv

        lines, texts = cls.extract_lines_and_texts(pdf_path, page_index)
        tables = cls.recover_tables(lines, texts)

        csv_path.parent.mkdir(parents=True, exist_ok=True)
        with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            if not tables or not tables[0].grid:
                for t in texts:
                    writer.writerow([t["text"]])
            else:
                for row in tables[0].grid:
                    writer.writerow(row)
        return csv_path.exists() and csv_path.stat().st_size > 0

    @classmethod
    def convert_pdf_to_html(cls, pdf_path: Path, html_path: Path, page_index: int = 0) -> bool:
        """PDF 표를 복원하여 웹 표준 HTML Table 문서로 변환 저장합니다."""
        import html

        lines, texts = cls.extract_lines_and_texts(pdf_path, page_index)
        tables = cls.recover_tables(lines, texts)

        html_path.parent.mkdir(parents=True, exist_ok=True)
        rows_html: List[str] = []
        if not tables or not tables[0].grid:
            for t in texts:
                rows_html.append(f"<tr><td>{html.escape(str(t['text']))}</td></tr>")
        else:
            for r_idx, row in enumerate(tables[0].grid):
                tag = "th" if r_idx == 0 else "td"
                cells = "".join(f"<{tag}>{html.escape(str(val))}</{tag}>" for val in row)
                rows_html.append(f"<tr>{cells}</tr>")

        doc_html = (
            "<!DOCTYPE html><html lang='ko'><head><meta charset='UTF-8'>"
            "<style>table{border-collapse:collapse;width:100%;font-family:'Malgun Gothic',sans-serif;font-size:13px;}"
            "th,td{border:1px solid #a0a0a0;padding:6px 10px;text-align:left;}th{background:#f2f4f8;}</style>"
            f"</head><body><table>{''.join(rows_html)}</table></body></html>"
        )
        html_path.write_text(doc_html, encoding="utf-8")
        return html_path.exists() and html_path.stat().st_size > 0

    @classmethod
    def convert_excel_to_csv(cls, excel_path: Path, csv_path: Path) -> bool:
        """XLSX 파일을 UTF-8 BOM CSV 파일로 변환합니다."""
        import csv

        csv_path.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.load_workbook(str(excel_path), data_only=True)
        ws = wb.active
        with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                writer.writerow(["" if v is None else str(v) for v in row])
        wb.close()
        return csv_path.exists() and csv_path.stat().st_size > 0

    @classmethod
    def convert_csv_to_excel(cls, csv_path: Path, excel_path: Path) -> bool:
        """CSV/TSV 파일을 서식이 적용된 XLSX 파일로 변환합니다."""
        import csv

        excel_path.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "변환된 표"
        delimiter = "\t" if csv_path.suffix.lower() == ".tsv" else ","
        with open(csv_path, "r", encoding="utf-8-sig", errors="ignore") as f:
            reader = csv.reader(f, delimiter=delimiter)
            for row in reader:
                ws.append(row)
        wb.save(str(excel_path))
        wb.close()
        return excel_path.exists() and excel_path.stat().st_size > 0
