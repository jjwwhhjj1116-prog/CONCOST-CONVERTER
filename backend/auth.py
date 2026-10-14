"""
사내 직원 인증 및 계정 관리 모듈 (Auth Module)
'CONNECT_TALK_직원계정.xlsx' 파일의 모든 시트를 순회하여 전체 임직원 계정을 로드하고,
초기 비밀번호 '1234' 부여 및 개인별 비밀번호 변경을 지원합니다.
"""
from dataclasses import asdict, dataclass
import hashlib
import json
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl

BASE_DIR: Path = Path(__file__).resolve().parent.parent
ACCOUNTS_EXCEL_PATHS: list[Path] = [
    BASE_DIR / "CONNECT_TALK_직원계정.xlsx",
    Path("C:/Users/user-owner/Downloads/CONNECT_TALK_직원계정.xlsx"),
    Path("e:/CONNECT_TALK_직원계정.xlsx"),
]

USERS_DB_PATH: Path = BASE_DIR / "storage" / "users_db.json"
USERS_DB_PATH.parent.mkdir(parents=True, exist_ok=True)


@dataclass
class EmployeeUser:
    """임직원 계정 데이터 클래스"""
    login_id: str
    name: str
    emp_no: str
    department: str
    rank: str
    email: str
    role: str
    password_hash: str  # SHA-256 해시


def hash_password(password: str) -> str:
    """비밀번호 SHA-256 해싱"""
    return hashlib.sha256(password.encode("utf-8")).hexdigest()


class AuthManager:
    """임직원 인증 및 계정 관리자"""

    def __init__(self) -> None:
        self.users: Dict[str, EmployeeUser] = {}
        self._load_or_initialize_users()

    def _find_excel_path(self) -> Optional[Path]:
        for p in ACCOUNTS_EXCEL_PATHS:
            if p.exists():
                return p
        return None

    def _load_or_initialize_users(self) -> None:
        """모든 시트를 순회하여 전 직원 계정 로드"""
        excel_path = self._find_excel_path()
        default_pwd_hash = hash_password("1234")

        # 기존 변경된 비밀번호 보존용 맵
        existing_hashes: Dict[str, str] = {}
        if USERS_DB_PATH.exists():
            try:
                with open(USERS_DB_PATH, "r", encoding="utf-8") as f:
                    old_data = json.load(f)
                    for uid, udata in old_data.items():
                        existing_hashes[uid] = udata.get("password_hash", default_pwd_hash)
            except Exception:
                pass

        if excel_path:
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            for sheetname in wb.sheetnames:
                ws = wb[sheetname]
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    continue

                headers = [str(h or "").strip() for h in rows[0]]
                if "로그인 ID" not in headers:
                    continue

                id_col = headers.index("로그인 ID")
                name_col = headers.index("이름") if "이름" in headers else 0
                no_col = headers.index("사번") if "사번" in headers else 2
                dept_col = headers.index("부서") if "부서" in headers else 3
                rank_col = headers.index("직급") if "직급" in headers else 4
                email_col = headers.index("이메일") if "이메일" in headers else 5
                role_col = headers.index("시스템 권한") if "시스템 권한" in headers else 7

                for r in rows[1:]:
                    if not r or len(r) <= id_col or not r[id_col]:
                        continue
                    lid = str(r[id_col]).strip().lower()
                    if lid in ("로그인 id", "id", "none", ""):
                        continue

                    name = str(r[name_col]).strip() if len(r) > name_col and r[name_col] else lid
                    emp_no = str(r[no_col] or "").strip() if len(r) > no_col else ""
                    dept = str(r[dept_col] or "").strip() if len(r) > dept_col else sheetname
                    rank = str(r[rank_col] or "").strip() if len(r) > rank_col else ""
                    email = str(r[email_col] or f"{lid}@concost.co.kr").strip() if len(r) > email_col else f"{lid}@concost.co.kr"
                    role = str(r[role_col] or "USER").strip() if len(r) > role_col else "USER"

                    # 기존 저장된 비밀번호가 있으면 유지, 없으면 1234
                    pwd_hash = existing_hashes.get(lid, default_pwd_hash)

                    self.users[lid] = EmployeeUser(
                        login_id=lid,
                        name=name,
                        emp_no=emp_no,
                        department=dept,
                        rank=rank,
                        email=email,
                        role=role,
                        password_hash=pwd_hash,
                    )
            wb.close()

        # 주요 관리자 계정 보장
        for name, lid, dept, rank, role in [
            ("유종욱", "yjw", "개발 TF", "실장", "SYSTEM_ADMIN"),
            ("박용진", "pyj", "개발 TF", "수석", "SYSTEM_ADMIN"),
        ]:
            if lid not in self.users:
                self.users[lid] = EmployeeUser(
                    login_id=lid,
                    name=name,
                    emp_no="1001",
                    department=dept,
                    rank=rank,
                    email=f"{lid}@concost.co.kr",
                    role=role,
                    password_hash=existing_hashes.get(lid, default_pwd_hash),
                )

        self._save_users()

    def _save_users(self) -> None:
        """사용자 DB를 JSON으로 영속 저장"""
        data = {uid: asdict(user) for uid, user in self.users.items()}
        with open(USERS_DB_PATH, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)

    def authenticate(self, login_id: str, password: str) -> Optional[EmployeeUser]:
        """로그인 검증"""
        lid = login_id.strip().lower()
        user = self.users.get(lid)
        if not user:
            return None
        if user.password_hash == hash_password(password):
            return user
        return None

    def change_password(self, login_id: str, old_pwd: str, new_pwd: str) -> bool:
        """개인설정 비밀번호 변경"""
        user = self.authenticate(login_id, old_pwd)
        if not user:
            return False
        user.password_hash = hash_password(new_pwd)
        self._save_users()
        return True

    def get_employee_public_list(self) -> List[Dict[str, Any]]:
        """로그인 드롭다운용 임직원 공개 목록"""
        result = []
        for u in self.users.values():
            result.append({
                "login_id": u.login_id,
                "name": u.name,
                "department": u.department,
                "rank": u.rank,
                "email": u.email,
                "role": u.role,
            })
        return sorted(result, key=lambda x: (x["department"], x["name"]))
